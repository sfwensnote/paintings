"""Join the V0.6 narrative JSON with the V0.6 anchor workbook."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
JSON_PATH = ROOT / "artwork_personality_assets_v1.json"
WORKBOOK_PATH = ROOT / "艺术人格测试_32作品评分与结果叙事_V0.6.xlsx"
OUTPUT_PATH = ROOT / "database" / "artworks.json"


def main() -> None:
    asset_file = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    assets = {item["type_id"]: item for item in asset_file["assets"]}
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    rows = workbook["32类型与锚点"].iter_rows(values_only=True)
    headers = next(rows)
    columns = {value: index for index, value in enumerate(headers) if value}
    required = ("类型ID", "编码(C-S-A-E-I)", "锚点作品", "作者", "年代", "文化/流派")
    missing = [field for field in required if field not in columns]
    if missing:
        raise ValueError(f"Workbook is missing columns: {', '.join(missing)}")

    records = []
    seen = set()
    for row in rows:
        type_id = row[columns["类型ID"]]
        if not type_id:
            continue
        type_id = str(type_id)
        if type_id in seen:
            raise ValueError(f"Duplicate type id in workbook: {type_id}")
        if type_id not in assets:
            raise ValueError(f"No JSON narrative data found for {type_id}")
        seen.add(type_id)
        asset = assets[type_id]
        title = str(row[columns["锚点作品"]]).strip()
        artist = str(row[columns["作者"]]).strip()
        type_code = str(row[columns["编码(C-S-A-E-I)"]]).strip()
        expected_code = "-".join("1" if asset["dimensions"][key] >= 50 else "0" for key in
                                 ("connection", "structure", "affect", "exploration", "imagination"))
        if type_code != expected_code or asset.get("type_code") != type_code:
            raise ValueError(f"Type encoding/vector mismatch for {type_id}: workbook={type_code}, scores={expected_code}")
        if title != asset["artwork"]:
            raise ValueError(f"Anchor mismatch for {type_id}: workbook={title}, JSON={asset['artwork']}")
        records.append({
            "artwork_id": f"A{int(type_id[1:]):03d}",
            "type_id": type_id,
            "type_code": type_code,
            "name_cn": title,
            "name_en": "",
            "artist_cn": artist if not artist.isascii() else "",
            "artist_en": artist if artist.isascii() else "",
            "year": str(row[columns["年代"]]).strip(),
            "culture_and_style": str(row[columns["文化/流派"]]).strip(),
            "dimensions": asset["dimensions"],
            "keywords": asset["keywords"],
            "archetype_name": asset["archetype_name"],
            "hit_line": asset["hit_line"],
            "life_narrative": asset["life_narrative"],
            "why_this_artwork": asset["why_this_artwork"],
            "art_context": f"{row[columns['年代']]} · {row[columns['文化/流派']]}。",
            "image_url": "",
            "image_source": "",
            "copyright_status": "图片来源与版权待复核",
            "status": "active",
            "catalog_kind": "formal_seed",
            "review_status": asset.get("review_status", "pending_human_review"),
            "version": "A_V0.6",
        })
    expected_ids = {f"T{i:02d}" for i in range(1, 33)}
    if seen != expected_ids or set(assets) != expected_ids:
        raise ValueError(f"Expected exact T01–T32 coverage. Workbook={len(seen)}, JSON={len(assets)}")
    OUTPUT_PATH.write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {len(records)} aligned V0.6 artwork records to {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
